"""Non-salary `.docx` / `.xlsx` content reader (Sprint 7b-1, ADR-0021).

A CONTENT-EXTRACTION UTILITY ONLY — it turns a docx/xlsx into structured text a
human (7b-1) or, later, the AI (7b-2) can read. It does NOT decide scope and does
NOT segment into facts (that is 7b-2). It writes NOTHING and never migrates
(ADR-0007): it reads the original from S3 and RETURNS content; hr-backend
persists it (as display `document_pages`, never `document_chunks` — these are
queried-not-embedded, ADR-0006).

This is the format extension to ADR-0010 (which is PDF-only): docx via
python-docx, xlsx via the salary path's openpyxl. It is deliberately separate
from `salary.py`, which hunts a salary GRID and discards everything else — a
reference source is non-salary and is routed here by its `document_type` tag, not
by content (Invariant 2). A salary `.xlsx` never reaches this reader.

Return envelope (canonical, one row per docx section / xlsx sheet — Q8):
    { format, pages: [{ page_number, label, text, locator }] }
hr-backend stores `pages` verbatim into `document_pages` for display reuse.
"""

from __future__ import annotations

import io

from .storage import get_object_bytes


def read_structured(storage_key: str, document_uuid: str, fmt: str) -> dict:
    """Read a docx/xlsx from S3 → { format, pages:[...] }. `fmt` ∈ docx | xlsx."""
    raw = get_object_bytes(storage_key)
    if fmt == "docx":
        return {"format": "docx", "pages": _read_docx(raw)}
    if fmt == "xlsx":
        return {"format": "xlsx", "pages": _read_xlsx(raw)}
    raise ValueError(f"unsupported reference format '{fmt}' (expected docx|xlsx)")


# --- docx --------------------------------------------------------------------

def _iter_block_items(parent):
    """Yield Paragraph and Table objects from a docx body IN DOCUMENT ORDER (the
    standard python-docx recipe — paragraphs and tables interleave)."""
    from docx.document import Document as _Document
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    parent_elm = parent.element.body if isinstance(parent, _Document) else parent._tc
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def _table_to_lines(table) -> list[str]:
    lines = []
    for row in table.rows:
        cells = [(" ".join(c.text.split())).strip() for c in row.cells]
        if any(cells):
            lines.append(" | ".join(cells))
    return lines


def _read_docx(raw: bytes) -> list[dict]:
    """Split a docx into sections at Heading-styled paragraphs (the multi-scope
    periodo docx use a province/section heading per block). If there are no
    headings, the whole document is one section."""
    from docx import Document as DocxDocument

    doc = DocxDocument(io.BytesIO(raw))

    sections: list[dict] = []
    current = {"heading": None, "lines": []}

    def flush():
        if current["lines"]:
            sections.append({"heading": current["heading"], "lines": list(current["lines"])})

    for block in _iter_block_items(doc):
        # Paragraph vs Table: a Paragraph has a `.style`; a Table has `.rows`.
        if hasattr(block, "rows"):
            current["lines"].extend(_table_to_lines(block))
            continue
        text = " ".join(block.text.split()).strip()
        if not text:
            continue
        style = (getattr(block.style, "name", "") or "")
        is_heading = style.lower().startswith("heading") or style.lower().startswith("título")
        if is_heading and current["lines"]:
            flush()
            current = {"heading": text, "lines": [text]}
        else:
            if is_heading:
                current["heading"] = text
            current["lines"].append(text)

    flush()

    if not sections:
        # No extractable content — surface an explicit empty page so the human
        # sees "(sin texto extraíble)" rather than a silent blank.
        return [{"page_number": 1, "label": "Documento", "text": "", "locator": "docx"}]

    pages = []
    for i, sec in enumerate(sections, start=1):
        label = sec["heading"] or f"Sección {i}"
        pages.append(
            {
                "page_number": i,
                "label": label,
                "text": "\n".join(sec["lines"]),
                "locator": f"section:{i}",
            }
        )
    return pages


# --- xlsx --------------------------------------------------------------------

def _read_xlsx(raw: bytes) -> list[dict]:
    """One page per sheet — the cells rendered as a readable text grid. Reuses
    openpyxl (already a salary dependency). This does NOT parse a salary grid; it
    surfaces the sheet content verbatim for a human to read (the salary path is a
    different, deliberately-tagged document — Invariant 2)."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    pages = []
    for idx, name in enumerate(wb.sheetnames, start=1):
        ws = wb[name]
        lines = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).strip() for v in row]
            if any(cells):
                lines.append(" | ".join(cells))
        pages.append(
            {
                "page_number": idx,
                "label": name,
                "text": "\n".join(lines),
                "locator": f"sheet:{name}",
            }
        )
    if not pages:
        return [{"page_number": 1, "label": "Hoja", "text": "", "locator": "sheet:1"}]
    return pages
